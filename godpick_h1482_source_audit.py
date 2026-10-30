# -*- coding: utf-8 -*-
"""H148.2: same-run source readiness and rejection provenance (research only).

Audits supplied current-run data. Never queries markets, invents OHLCV, changes
H109/H113/H126 authority, or equates a watch-list with a trading suggestion.
"""
from __future__ import annotations

from collections import Counter
from typing import Any, Mapping

import pandas as pd

from godpick_h147_five_gate_review import _code, _day, _n, _s
from godpick_h148_strength_provenance import _evidence_time

VERSION = 'V191-H148.2-SOURCE-READINESS-AND-REJECTION-TRUTH'
DETAIL_COLUMNS = (
    '股票代號', '股票名稱', '流程位置', '研究資料日', 'H114證據日', 'H114資料狀態',
    'H114有效根數', 'H114量比', 'H113流動性', 'H109風控', 'H123法人',
    'H124大戶', 'H1481研究分層', '完整隔離原因', '最優先修復項', '交易權限',
)


def _frame(obj: Any) -> pd.DataFrame:
    return obj if isinstance(obj, pd.DataFrame) else pd.DataFrame()


def _rows(frame: pd.DataFrame) -> tuple[dict[str, dict], Counter, list[str]]:
    found: dict[str, dict] = {}
    duplicates: Counter = Counter()
    invalid: list[str] = []
    if frame.empty:
        return found, duplicates, invalid
    for i, row in enumerate(frame.to_dict('records')):
        raw = row.get('股票代號', '')
        code = _code(raw)
        if not code:
            invalid.append(f'第{i+1}列:{_s(raw) or "缺代號"}'[:64])
        elif code in found:
            duplicates[code] += 1
        else:
            found[code] = row
    return found, duplicates, invalid


def _priority(why: str) -> str:
    # Order by actionable ROOT CAUSE, not the first arbitrary error listed.
    for token, label in (
        ('同代號', '先核對同輪重複與來源衝突'),
        ('交易日', '補齊正確交易日K線與資料來源'),
        ('K線', '補齊H114有效OHLCV與日期'),
        ('H113', '核對成交額、20日均額及張數'),
        ('成交額', '核對成交額與張數單位'),
        ('量比', '核對本日成交量與前20日均量'),
        ('H108', '補H108族群廣度、資金流及跨期日期'),
        ('H109', '核對H109原否決與前後期依據'),
        ('均線', '補MA20/MA60與乖離證據'),
        ('H123', '補齊官方法人逐日來源'),
        ('H124', '補齊TDCC兩週來源'),
    ):
        if token in why:
            return label
    return '逐股核對當輪資料與H126決策'


def build_source_audit(*, candidate_frame: Any, audit_frame: Any, trace_frame: Any,
                       quarantine_frame: Any, research_day: Any,
                       trace_meta: Mapping[str, Any] | None = None
                       ) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    """Return (repair summary, per-audit-stock rejection details, diagnostics).

    Candidate rows with no H113 evidence may be located but must NEVER be given
    pseudo-valid H114/volume/sector evidence by joining another stock or date.
    """
    candidate = _frame(candidate_frame)
    audit = _frame(audit_frame)
    trace = _frame(trace_frame)
    quarantine = _frame(quarantine_frame)
    cand, cdups, cbad = _rows(candidate)
    ad, adups, abad = _rows(audit)
    tr, tdups, tbad = _rows(trace)
    qa, qdups, qbad = _rows(quarantine)
    cutoff = _day(research_day)
    stage = Counter()
    rejection = Counter()
    source = Counter()
    details = []
    for code, row in ad.items():
        traced = tr.get(code, {})
        isolated = qa.get(code, {})
        found_in = _s(traced.get('H148漏網位置')) or 'H113稽核／H148追蹤未見'
        date_status = _evidence_time(row, cutoff)
        source[date_status] += 1
        h114 = _s(row.get('H114K線資料狀態')) or '缺H114狀態'
        st = _s(isolated.get('H1481研究分層')) or '非H1481隔離候選'
        why = _s(isolated.get('H1481入榜／隔離依據'))
        if code in adups or code in tdups or code in qdups:
            why = '同代號多筆來源／需隔離' + ('；'+why if why else '')
        if not why and date_status != '同日完整OHLCV（內部標記）':
            why = '缺同研究交易日有效K線：'+date_status
        for reason in [x.strip() for x in why.split('；') if x.strip()]:
            rejection[reason] += 1
        stage[found_in] += 1
        details.append({
            '股票代號':code, '股票名稱':_s(row.get('股票名稱')),
            '流程位置':found_in, '研究資料日':str(cutoff) if cutoff else '未核實',
            'H114證據日':_s(row.get('H114K線證據日')) or '未提供',
            'H114資料狀態':h114,'H114有效根數':_n(row,'H114K線有效根數'),
            'H114量比':_n(row,'H114K線量比'),
            'H113流動性':_s(row.get('H113流動性合格')) or '未證',
            'H109風控':_s(row.get('H109決策權威')) or '未知',
            'H123法人':_s(row.get('H123法人實證狀態')) or '待補證',
            'H124大戶':_s(row.get('H124大戶實證狀態')) or '待補證',
            'H1481研究分層':st, '完整隔離原因':why or '未隔離／研究資格仍以H126原表為準',
            '最優先修復項':_priority(why or date_status), '交易權限':'研究稽核／非買進',
        })
    # The 120-row H113 audit is a BOUNDED downstream sample. Capture diagnosis
    # candidates outside it without describing them as rejected by H113.
    # Put candidates with ACTUAL dated H114 proof first in this research-only
    # appendix; no proxy price, no fabricated H113 qualification or buy rank.
    candidate_outside = []
    for code in sorted(set(cand) - set(ad)):
        row = cand[code]
        traced = tr.get(code, {})
        state = _evidence_time(row, cutoff)
        valid = state == '同日完整OHLCV（內部標記）'
        ratio = _n(row, 'H114K線量比')
        candidate_outside.append((0 if valid and ratio is not None and ratio >= 1.2 else 1 if valid else 2,
           code, {
              '股票代號':code, '股票名稱':_s(row.get('股票名稱')),
              '流程位置':_s(traced.get('H148漏網位置')) or '診斷候選未進120檔H113樣本',
              '研究資料日':str(cutoff) if cutoff else '未核實',
              'H114證據日':_s(row.get('H114K線證據日')) or '未提供',
              'H114資料狀態':_s(row.get('H114K線資料狀態')) or '缺H114狀態',
              'H114有效根數':_n(row,'H114K線有效根數'),
              'H114量比':ratio,
              'H113流動性':'未進H113稽核樣本／不得推論合格',
              'H109風控':_s(row.get('H109決策權威')) or '未核',
              'H123法人':_s(row.get('H123法人實證狀態')) or '待補證',
              'H124大戶':_s(row.get('H124大戶實證狀態')) or '待補證',
              'H1481研究分層':'診斷池線索／未獲H113稽核',
              '完整隔離原因':'未進H113限量稽核樣本；'+(
                '已見同日內部K線與量比，仍需正式稽核' if valid and ratio is not None and ratio >= 1.2
                else 'H114量價尚未同日確認；'+state),
              '最優先修復項':'核對H113選樣排序與入池原因' if valid else _priority(state),
              '交易權限':'診斷線索／非買進',
            }))
    candidate_outside.sort(key=lambda x:(x[0],x[1]))
    details.extend(row for _,_,row in candidate_outside)
    detail = pd.DataFrame(details, columns=DETAIL_COLUMNS)
    missing = sorted(set(cand) - set(ad))
    outside = sorted(set(ad) - set(cand))
    without_trace = sorted(set(cand) - set(tr))
    report = trace_meta if isinstance(trace_meta, Mapping) else {}
    repair = [
        ('診斷股票原始列數', len(candidate), '不等同全市場有效股票數'),
        ('診斷有效四碼唯一股票', len(cand), '以4位股票代號匹配，無效代號不得無聲刪去'),
        ('診斷無效或缺代號列數', len(cbad), '核實商品種類/代號：' + '、'.join(cbad[:6])),
        ('診斷重複代號數', len(cdups), '需查同輪資料是否混合'),
        ('H113稽核原始列數', len(audit), '此池並非全市場'),
        ('H113有效四碼唯一股票', len(ad), '不可拿120檔外推全市場'),
        ('診斷有但未進H113', len(missing), 'H113為限量稽核樣本；不代表真正漏掃或符合交易條件'),
        ('H113外已有同日放量K證據的診斷線索', sum(tier==0 for tier,_,_ in candidate_outside), '研究線索仍需H113逐項核實與原H126授權'),
        ('H113不在診斷池', len(outside), '核對是否跨輪或來源矛盾'),
        ('診斷股票未見於H148追蹤', len(without_trace), '逐股找缺失代號與處理原因'),
        ('H148追蹤有效四碼唯一股票', len(tr), '與診斷/H113總和應按聯集比對'),
        ('H114同研究日完整K線檔數', source['同日完整OHLCV（內部標記）'], '僅內部標記，需比對官方來源'),
        ('H114缺交易日檔數', source['缺H114交易日'], '檢查每日歷史K線分析是否進入掃描列'),
        ('H114不足60根檔數', sum(n for k,n in source.items() if '不足60' in k), '補足歷史資料，不可假補'),
        ('H1481隔離有代號檔數', len(qa), '全部保留原因，非正式推薦'),
        ('同輪診斷母體已核', '是' if report.get('proven_diagnosis_pool') is True else '否', '不代表已證實全市場市場覆蓋'),
    ]
    repair_df = pd.DataFrame(repair, columns=['檢查項目','結果','說明'])
    meta = {
        'version':VERSION, 'research_day':str(cutoff) if cutoff else '未核實',
        'candidate_rows':len(candidate), 'candidate_unique':len(cand),
        'candidate_invalid':len(cbad), 'candidate_invalid_examples':cbad[:12],
        'candidate_duplicate_codes':sorted(cdups), 'audit_unique':len(ad),
        'trace_unique':len(tr), 'trace_invalid':len(tbad),
        'audit_invalid':len(abad), 'quarantine_unique':len(qa),
        'candidate_not_h113':len(missing), 'not_h113_examples':missing[:20],
        'candidate_with_dated_k_outside_h113':sum(tier==0 for tier,_,_ in candidate_outside),
        'detail_rows':len(detail),
        'h113_not_candidate':len(outside), 'h113_extra_examples':outside[:20],
        'untraced_candidate_count':len(without_trace), 'untraced_examples':without_trace[:20],
        'same_day_h114':source['同日完整OHLCV（內部標記）'],
        'source_status_counts':dict(source),
        'rejection_reason_counts':dict(rejection),
        'stage_counts':dict(stage),
        'scope_proven':bool(report.get('proven_diagnosis_pool') is True and not cbad and not cdups),
        'h126_modified':False,
    }
    return repair_df, detail, meta


__all__ = ['VERSION','DETAIL_COLUMNS','build_source_audit']


def append_source_audit_to_sheet(ws: Any, repair: Any, details: Any, *,
                                 max_rows: int = 500) -> dict[str, int]:
    """Append an explicitly NON-RECOMMENDATION evidence appendix on sheet #1.

    Keeps two-workbook tabs, the H126 authority sheet and the top-12 table
    structurally untouched. The existing top-table filter is not expanded.
    """
    from openpyxl.styles import Font, PatternFill, Alignment
    repair, details = _frame(repair), _frame(details)
    first = max(ws.max_row + 3, 11)
    ws.cell(first, 1, 'H148.2｜資料來源、強勢股漏網與隔離查核（以下全部非推薦／非交易）')
    ws.merge_cells(start_row=first, start_column=1, end_row=first, end_column=15)
    ws.cell(first, 1).fill = PatternFill('solid', fgColor='12344A')
    ws.cell(first, 1).font = Font(bold=True, color='FFFFFF', size=13)
    ws.row_dimensions[first].height = 29
    at = first+1
    for _, record in repair.head(30).iterrows():
        ws.cell(at,1,str(record.get('檢查項目','')))
        ws.cell(at,2,record.get('結果',''))
        ws.cell(at,3,str(record.get('說明',''))[:900])
        ws.cell(at,3).alignment = Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[at].height=29
        at+=1
    at+=2
    headers=['股票代號','股票名稱','流程位置','H114證據日','H114資料狀態',
             'H114有效根數','H114量比','H113流動性','H109風控','H123法人',
             'H124大戶','完整隔離原因','最優先修復項','交易權限']
    for col,header in enumerate(headers,1):
        cell=ws.cell(at,col,header)
        cell.fill=PatternFill('solid',fgColor='6B2735')
        cell.font=Font(color='FFFFFF',bold=True)
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[at].height=29
    source_map={
        '股票代號':'股票代號','股票名稱':'股票名稱','流程位置':'流程位置',
        'H114證據日':'H114證據日','H114資料狀態':'H114資料狀態',
        'H114有效根數':'H114有效根數','H114量比':'H114量比',
        'H113流動性':'H113流動性','H109風控':'H109風控',
        'H123法人':'H123法人','H124大戶':'H124大戶',
        '完整隔離原因':'完整隔離原因','最優先修復項':'最優先修復項',
        '交易權限':'交易權限',
    }
    listed=0
    for rec in details.head(max(0,int(max_rows))).to_dict('records'):
        at+=1
        for col,key in enumerate(headers,1):
            val=rec.get(source_map[key])
            # Prevent user-controlled spreadsheet formula / hyperlink injection.
            if isinstance(val,str) and val.lstrip().startswith(('=','+','-','@')):
                val="'"+val
            cell=ws.cell(at,col,val)
            if col in (3,12,13):
                cell.alignment=Alignment(vertical='center',wrap_text=True)
        ws.row_dimensions[at].height=45
        listed+=1
    return {'summary':min(30,len(repair)),'details':listed,'first_row':first}


__all__.append('append_source_audit_to_sheet')
