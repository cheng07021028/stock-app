# -*- coding: utf-8 -*-
"""Execute the verbatim Page07 Excel writer (not a mocked reimplementation)."""
import ast
import io
import json
from pathlib import Path
import pandas as pd
from godpick_h144_focus_board import build_focus_board

ROOT=Path(__file__).resolve().parents[1]
PAGE=ROOT/'pages'/'7_股神推薦.py'


def _actual_writer():
    source=PAGE.read_text('utf-8')
    tree=ast.parse(source)
    funcs=['_excel_safe_value','_excel_safe_df','_write_df_to_ws_fast_h84']
    found=[x for x in tree.body if isinstance(x,ast.FunctionDef) and x.name in funcs]
    assert len(found)==3
    begin='    sheets = [\n        ("00_主管重點觀察榜", h144_focus,'
    assert source.count(begin)==1
    start=source.index(begin)
    end=source.index('    return output.getvalue()',start)+len('    return output.getvalue()')
    exact=source[start:end]
    wrapper='''def render(h144_focus,h144_focus_meta,auth,bridge):
    from openpyxl import Workbook
    import io
    wb=Workbook()
    wb.remove(wb.active)
    output=io.BytesIO()
    tables={'h126_authority':auth}
    _h145_bridge={'status':bridge}
'''+exact+'\n'
    ctx={'pd':pd,'json':json}
    exec(compile(ast.fix_missing_locations(ast.Module(body=found,type_ignores=[])),str(PAGE),'exec'),ctx,ctx)
    exec(compile(wrapper,str(PAGE),'exec'),ctx,ctx)
    return ctx['render']


def test_verbatim_page07_writer_round_trip_twelve_ties():
    from openpyxl import load_workbook
    ids=['1326','1717','2115','2317','2337','2344','2376','2395','3037','3044','3680','3702']
    audit=pd.DataFrame([{'股票代號':c,'股票名稱':'測試'+c,'類別':'測試',
       'H110最終研究優先分':52,'H109決策權威':'FADE_VETO',
       'H123法人實證狀態':'PENDING_EVIDENCE','H124大戶實證狀態':'PENDING_EVIDENCE'} for c in ids])
    authority=pd.DataFrame([{'H126順位':None,'股票代號':None,
      '股票名稱':'本輪沒有合格個股','H126能否買進':'禁止買進'}])
    board,meta=build_focus_board({'audit':audit,'h126_authority':authority},
      {'status':'EXPORT_SCOPE_INCOMPLETE','scope_incomplete':True,
       'scan_count':0,'diagnosis_count':0,'reported_final':8,'export_count':1667})
    blob=_actual_writer()(board,meta,authority,'BLOCKED_PARTIAL_EXPORT_NO_TRUSTED_RECOVERY')
    assert len(blob)>1500
    book=load_workbook(io.BytesIO(blob),read_only=False)
    assert book.sheetnames==['00_主管重點觀察榜','00_H126主管決策']
    ws=book.worksheets[0]
    assert ws.freeze_panes=='D8'
    assert ws.auto_filter.ref==f'A7:Q{ws.max_row}'
    assert ws['A7'].value=='觀察順位'
    assert ws['B7'].value=='股票代號'
    assert ws['A8'].value=='並列／無法區分'
    assert all(ws.cell(row,1).value=='並列／無法區分' for row in range(8,20))
    assert all(ws.cell(row,5).value==52 for row in range(8,20))
    assert '掃描數為0' in ws['B3'].value
    assert '全部有效研究分相同' in ws['B6'].value
    assert 'H126實際個股授權 0' in ws['B5'].value
    assert ws['M8'].value == '未列入H126合格決策'
    assert ws['P8'].value == '禁止依本榜買進'
    assert len(ws.merged_cells.ranges)==6
    auth=book.worksheets[1]
    assert auth['A1'].value=='H126順位'
    assert auth['C2'].value=='本輪沒有合格個股'
    book.close()


def test_verbatim_page07_writer_diverse_ranks_and_conflicts():
    from openpyxl import load_workbook
    audit=pd.DataFrame([
      {'股票代號':'2317','H110最終研究優先分':75},
      {'股票代號':'2344','H110最終研究優先分':75},
      {'股票代號':'3037','H110最終研究優先分':61},
      {'股票代號':'3680','H110最終研究優先分':48},
    ])
    authority=pd.DataFrame([{'股票代號':c,'H126能否買進':'否'} for c in ('2317','2344','3037','3680')])
    board,meta=build_focus_board({'audit':audit,'h126_authority':authority},
        {'status':'SCOPE_RECONCILED_FOR_DISPLAY','scan_count':4,
        'diagnosis_count':4,'reported_final':4,'export_count':4})
    book=load_workbook(io.BytesIO(_actual_writer()(board,meta,authority,'READY')),read_only=False)
    ws=book.worksheets[0]
    assert [ws.cell(i,1).value for i in (8,9,10,11)]==['並列第1','並列第1',3,4]
    assert ws.auto_filter.ref=='A7:Q11'
    book.close()
