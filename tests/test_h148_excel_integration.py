# -*- coding: utf-8 -*-
"""Render the ACTUAL modified page-7 Excel body (AST extracted) to xlsx."""
import ast
import io
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from godpick_h147_five_gate_review import build_five_gate_review
from godpick_h144_focus_board import build_focus_board
from godpick_h148_strength_provenance import build_strength_trace

ROOT=Path(__file__).resolve().parents[1]


def _render():
    page=ROOT/'pages'/'7_股神推薦.py'
    content=page.read_text('utf-8')
    tree=ast.parse(content)
    names=['_excel_safe_value','_excel_safe_df','_write_df_to_ws_fast_h84']
    funcs=[f for f in tree.body if isinstance(f,ast.FunctionDef) and f.name in names]
    assert len(funcs)==len(names)
    beg=content.index('    sheets = [\n        ("00_主管重點觀察榜", h144_focus,')
    end=content.index('    return output.getvalue()',beg)+len('    return output.getvalue()')
    body='''def render(board,meta,authority,five,trace,summary):
    from openpyxl import Workbook
    import io
    wb=Workbook()
    wb.remove(wb.active)
    output=io.BytesIO()
    h144_focus=board
    h144_focus_meta=meta
    tables={'h126_authority':authority}
    h147_five_gates=five
    h147_five_meta={'reviewed':len(five),'data_gaps':0,'already_ignited':0}
    h148_trace=trace
    h148_meta=summary
    _h148_export_meta=summary
    _h147_scan_diag={'status':'SCHEMA_RECONCILED'}
    _h145_bridge={'status':'SUPPLIED_FULL_SCAN_MATCHED'}
'''+content[beg:end]+'\n'
    ns={'pd':pd}
    exec(compile(ast.fix_missing_locations(ast.Module(body=funcs,type_ignores=[])),str(page),'exec'),ns,ns)
    exec(compile(body,str(page),'exec'),ns,ns)
    return ns['render']


def test_h148_real_two_sheet_excel_contains_h147_and_h148_without_trade_override():
    raw={'股票代號':'2317','股票名稱':'測試','類別':'AI設備','H110最終研究優先分':52,
         'H110研究資訊截止日':'2026-10-28','H114K線證據日':'2026-10-28',
         'H114K線資料狀態':'VALID_OHLCV','H114K線有效根數':80,
         'H113當日成交額百萬':310,'H113二十日均成交額百萬':280,
         'H113當日成交量張':1500,'H113流動性合格':'是','H114K線量比':1.5,
         'H114距20日高點%':-1.0,'收盤距MA20%':4,'收盤距MA60%':9,
         'H105發動階段':'BH1','H108族群生命週期':'IGNITION',
         'H109決策權威':'RESEARCH_ELIGIBLE','H123法人實證狀態':'PENDING_EVIDENCE',
         'H124大戶實證狀態':'PENDING_EVIDENCE'}
    audit=pd.DataFrame([raw]);authority=pd.DataFrame([{'股票代號':'2317','H126能否買進':'否'}])
    state={'status':'SCOPE_RECONCILED_FOR_DISPLAY','scope_incomplete':False,
           'scan_count':10,'diagnosis_count':1,'reported_final':1,'export_count':1}
    board,quality=build_focus_board({'audit':audit,'h126_authority':authority},state)
    five,_=build_five_gate_review(audit,research_day='2026-10-28')
    trace,_,meta=build_strength_trace(candidate_frame=audit,audit_frame=audit,result_frame=audit,
        authority_frame=authority,scan_report={'total_count':10,'candidate_diagnosis_count':1,'final_candidate_count':1},
        h145_bridge={'candidate_complete':True,'same_run_proven':True},research_day='2026-10-28')
    book=load_workbook(io.BytesIO(_render()(board,quality,authority,five,trace,meta)))
    assert book.sheetnames==['00_主管重點觀察榜','00_H126主管決策']
    first,formal=book.worksheets
    assert first['R7'].value=='H147｜成交量能'
    assert first['X7'].value=='H147｜行情日期'
    assert first['Y7'].value=='H148研究分類'
    assert first['AF7'].value=='H148補證優先'
    assert first['Y8'].value=='突破確認研究／非買點'
    assert '診斷候選' in first['Z8'].value or '最終作戰' in first['Z8'].value
    assert first['AB8'].value=='突破'
    assert 'H148查核1檔' in first['B6'].value
    assert first['P8'].value=='禁止依本榜買進'
    assert first.freeze_panes=='D8'
    assert formal['G2'].value in (None,'') or first['P8'].value=='禁止依本榜買進'
    assert '否' in [c.value for c in formal[2]]
    book.close()
