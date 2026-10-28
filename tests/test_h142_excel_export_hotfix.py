# -*- coding: utf-8 -*-
"""Regression: H142 appended h141_quality before manager_quality_gate initialized it.

Runs the VERBATIM manager-export block from Page07, not a reimplemented copy.
No Streamlit, network, GitHub, stock selection or persistence writes.
"""
from __future__ import annotations

import io
import sys
import ast
import json
import unittest
from pathlib import Path
from textwrap import dedent

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def _export_block(tables):
    page = (ROOT / 'pages' / '7_股神推薦.py').read_text(encoding='utf-8')
    start_anchor = '    h129_risk_watch = _h129_risk_watch(tables)\n'
    end_anchor = '    _h126_manager_tabs = {\n'
    assert page.count(start_anchor) == 1, 'Manager export code block moved or duplicated'
    start = page.index(start_anchor)
    end = page.index(end_anchor, start)
    block = dedent(page[start:end])
    # Inputs normally prepared by earlier canonical H126/H129 stages.
    env = {
        'pd': pd,
        'tables': tables,
        'health': pd.DataFrame([{'項目': 'smoke', '數值': 'ok'}]),
        'h129_t1_health': pd.DataFrame([{'項目':'h129', '數值':'ok'}]),
        'h102_t1_review': pd.DataFrame(),
        'sector': pd.DataFrame(),
        '_h129_risk_watch': lambda frames: pd.DataFrame([{'狀態': '風控有效'}]),
    }
    exec(compile(block, str(ROOT / 'pages' / '7_股神推薦.py'), 'exec'), env, env)
    return env


def _read_workbook(sheets):
    # Compile the ACTUAL Page07 writer and its sanitizers; no copied implementation.
    filename = str(ROOT / 'pages' / '7_股神推薦.py')
    tree = ast.parse(Path(filename).read_text(encoding='utf-8'), filename=filename)
    funcs = {'_excel_safe_value', '_excel_safe_df', '_write_df_to_ws_fast_h84'}
    nodes = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in funcs]
    assert len(nodes) == len(funcs)
    module = ast.fix_missing_locations(ast.Module(body=nodes, type_ignores=[]))
    ctx = {'pd': pd, 'json': json}
    exec(compile(module, filename, 'exec'), ctx, ctx)
    writer = ctx['_write_df_to_ws_fast_h84']
    book = Workbook()
    book.remove(book.active)
    for name, frame, fallback in sheets:
        writer(book, name, frame, fallback)
    buff = io.BytesIO()
    book.save(buff)
    buff.seek(0)
    return load_workbook(buff, read_only=True)


class H142ManagerExcelHotfix(unittest.TestCase):
    @staticmethod
    def _base(audit):
        return {
            'audit': pd.DataFrame(audit),
            'h126_authority': pd.DataFrame([
                {'股票代號': '4938', 'H126權威層別': 'B｜待驗證觀察（非買進）',
                 'H126能否買進': '否', 'H126認證資格': 'PENDING_EVIDENCE'},
                {'股票代號': '2324', 'H126權威層別': 'B｜待驗證觀察（非買進）',
                 'H126能否買進': '否', 'H126認證資格': 'PENDING_EVIDENCE'},
            ]),
            'h126_health': pd.DataFrame([{'項目': '官方來源', '數值':'PENDING_EVIDENCE'}]),
            'priority_overview': pd.DataFrame([{
                '股票代號': '2317','參考進場價': 200,
                '停損價': 180, '目標價': 260, 'RR': 3.0,
            }]),
            'actionable': pd.DataFrame(),
        }

    @staticmethod
    def _audit(n=120):
        return [{
            '股票代號':str(2000+i), '股票名稱':f'測試{i}', '類別':'AI伺服器' if i%2 else '化工',
            'H109決策權威':'FADE_VETO' if i < n - 2 else 'RESEARCH_ELIGIBLE',
            'H108族群生命週期':'FADE｜退潮' if i < n - 2 else 'PREHEAT',
            'H110研究資訊截止日':'2026-10-21',
            'H110證據狀態':'PARTIAL',
            'H123法人實證狀態':'PENDING_EVIDENCE',
            'H124大戶實證狀態':'PENDING_EVIDENCE',
        } for i in range(n)]

    def test_h141_quality_initialized_before_h142_append(self):
        result = _export_block(self._base(self._audit()))
        quality = result['h141_quality']
        labels = quality['項目'].astype(str).tolist()
        self.assertIn('H141匯出檢查版本', labels)
        self.assertIn('H109 FADE_VETO檔數', labels)
        self.assertEqual(dict(zip(quality['項目'], quality['數值']))['H109 FADE_VETO檔數'], 118)
        self.assertEqual(len(result['h142_evidence']), 120)
        self.assertEqual(len(result['tables']['actionable']), 0)

    def test_legacy_price_fields_stay_redacted(self):
        res = _export_block(self._base(self._audit(2)))
        old = res['h141_safe_legacy'].iloc[0]
        for price in ('參考進場價','停損價','目標價','RR'):
            self.assertTrue(pd.isna(old[price]), price)
        self.assertIn('不可買進', old['H141研究授權警示'])

    def test_manager_workbook_sheets_are_writable_and_readable(self):
        res = _export_block(self._base(self._audit()))
        workbook = _read_workbook(res['sheets'])
        # H144+ intentionally exports only two manager-facing sheets.
        # H141/H142 remain computed internally and must be checked above,
        # not reintroduced as separate manager workbook tabs.
        self.assertEqual(workbook.sheetnames, ['00_主管重點觀察榜', '00_H126主管決策'])
        self.assertGreater(len(res['h141_quality']), 10)
        self.assertEqual(len(res['h142_evidence']), 120)
        self.assertIn('H141匯出檢查版本', res['h141_quality']['項目'].tolist())
        self.assertEqual(workbook['00_H126主管決策'].max_row, 3)
        workbook.close()

    def test_empty_candidates_still_export_without_local_variable_error(self):
        frames = self._base([])
        frames['h126_authority'] = pd.DataFrame()
        res = _export_block(frames)
        book = _read_workbook(res['sheets'])
        self.assertEqual(book.sheetnames, ['00_主管重點觀察榜', '00_H126主管決策'])
        self.assertIn('H141匯出檢查版本', res['h141_quality']['項目'].tolist())
        self.assertTrue(res['h142_evidence'].empty)
        book.close()

if __name__ == '__main__':
    unittest.main()
