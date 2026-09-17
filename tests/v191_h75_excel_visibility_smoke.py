# -*- coding: utf-8 -*-
import ast, json, tempfile
from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT=Path(__file__).resolve().parents[1]
page=(ROOT/'pages'/'7_股神推薦.py').read_text(encoding='utf-8')
tree=ast.parse(page)
need={'_excel_safe_value','_excel_safe_df','_write_df_to_ws'}
nodes=[n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in need]
assert {n.name for n in nodes}==need
mod=ast.Module(body=nodes, type_ignores=[])
ns={'pd':pd,'json':json}
exec(compile(ast.fix_missing_locations(mod), str(ROOT/'pages'/'7_股神推薦.py'), 'exec'), ns)
write=ns['_write_df_to_ws']

wb=Workbook(); wb.remove(wb.active)
main=pd.DataFrame([{
 '今日順位':1,'股票代號':'2330','股票名稱':'測試','H75主管決策層級':'W2｜優先觀察','H75今日結論':'優先觀察｜資金加速（非Formal）',
 'H75Formal接近度':77.2,'H75Formal主要缺口':'缺主流共振；上游權威=A-MINUS','H75主管一句話':'非Formal；等待主流與盤前重驗'
}])
for name in ['01_今日AI推薦','02_今日主流資金','03_推薦證據','04_推薦績效','05_系統健康']:
    write(wb,name,main,'無資料')
for name in ['H74新鮮主流資金','股神推薦總排名','T+1實戰真相']:
    write(wb,name,main,'無資料')
visible={'01_今日AI推薦','02_今日主流資金','03_推薦證據','04_推薦績效','05_系統健康'}
for ws in wb.worksheets:
    ws.sheet_state='visible' if ws.title in visible else 'hidden'
wb.active=wb.sheetnames.index('01_今日AI推薦')
path=Path(tempfile.gettempdir())/'h75_excel_visibility_smoke.xlsx'
wb.save(path)
check=load_workbook(path, read_only=False, data_only=False)
assert check.active.title=='01_今日AI推薦'
assert [ws.title for ws in check.worksheets if ws.sheet_state=='visible']==['01_今日AI推薦','02_今日主流資金','03_推薦證據','04_推薦績效','05_系統健康']
assert all(check[n].sheet_state=='hidden' for n in ['H74新鮮主流資金','股神推薦總排名','T+1實戰真相'])
assert check['01_今日AI推薦']['A1'].value=='今日順位'
print('PASS H75 Excel write/visibility')
