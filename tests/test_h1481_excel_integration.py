# -*- coding: utf-8 -*-
"""Verify the real H148.1 page-7 workbook body remains exactly two sheets."""
import io
import pandas as pd
from openpyxl import load_workbook
from test_h148_excel_integration import _render
from test_h1481_quality_gate import row, run, DAY
from godpick_h147_five_gate_review import build_five_gate_review


def test_h1481_real_excel_strong_primary_and_H126_untouched():
    strong=row('3035')
    veto=row('2344',H109決策權威='FADE_VETO')
    focus,quarantine,quality,trace,legacy=run([veto,strong])
    assert list(focus['股票代號'])==['3035'] and len(quarantine)==1
    audit=pd.DataFrame([veto,strong]); five,_=build_five_gate_review(audit,research_day=DAY)
    authority=pd.DataFrame([{'股票代號':'3035','H126能否買進':'否'},
                            {'股票代號':'2344','H126能否買進':'否'}])
    report={'review_unique':2,'scope_status':'已證明本輪診斷候選池',
            'zero_formal_diagnosis':'測試：無正式授權'}
    meta={'scanned':10,'diagnosed':2,'final':2,'received':2,'audit_count':2,
          'authority_count':2,'all_scores_tied':True,'scope_issues':[],
          'valid_score_rows':2,'score_distinct':1,'h126_buyable':0,'warning':'研究觀察資料'}
    out=_render()(focus,meta,authority,five,trace,report)
    wb=load_workbook(io.BytesIO(out),read_only=False)
    assert wb.sheetnames==['00_主管重點觀察榜','00_H126主管決策']
    ws,formal=wb.worksheets
    assert ws['B8'].value=='3035'
    assert ws['P8'].value=='禁止依本榜買進'
    assert ws['R7'].value=='H147｜成交量能'
    assert ws['Y7'].value=='H148研究分類'
    assert ws['AG7'].value=='H1481研究分層'
    assert ws['AJ7'].value=='H1481入榜／隔離依據'
    assert '同日K線60根以上' in ws['AH8'].value
    assert 'H108' in ws['AI8'].value
    assert '突破' in ws['AJ8'].value
    assert ws['B9'].value is None
    assert ws.freeze_panes=='D8'
    assert list(formal.values)[1][0] is not None
    assert '否' in [str(c.value) for r in formal for c in r]
    wb.close()


def test_h1481_empty_board_never_refills_score_ties_in_actual_xlsx():
    bad=row('2344',H109決策權威='FADE_VETO',H114K線證據日=None)
    board,blocked,info,trace,_=run([bad])
    assert board.empty and len(blocked)==1
    auth=pd.DataFrame([{'股票代號':'2344','H126能否買進':'否'}])
    meta={'scanned':10,'diagnosed':1,'final':1,'received':1,'audit_count':1,
          'authority_count':1,'all_scores_tied':False,'scope_issues':[],
          'valid_score_rows':1,'score_distinct':1,'h126_buyable':0,'warning':'僅供研究'}
    report={'review_unique':1,'scope_status':'局部快照','zero_formal_diagnosis':'缺完整證據'}
    blob=_render()(board,meta,auth,pd.DataFrame(),trace,report)
    wb=load_workbook(io.BytesIO(blob));ws=wb.worksheets[0]
    assert wb.sheetnames==['00_主管重點觀察榜','00_H126主管決策']
    assert 'H148.1' in ws['A1'].value
    assert ws['B8'].value not in ('2344',2344)
    assert 'H1481' in ws['AG7'].value
    wb.close()
