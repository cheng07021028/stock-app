# -*- coding: utf-8 -*-
"""The ACTUAL page-07 writer renders H147 R:X columns and keeps H146 A:Q."""
import ast
import io
from pathlib import Path
import pandas as pd
from godpick_h147_five_gate_review import build_five_gate_review
from godpick_h144_focus_board import build_focus_board

ROOT=Path(__file__).resolve().parents[1]
PAGE=ROOT/'pages'/'7_股神推薦.py'


def _writer():
    source=PAGE.read_text('utf-8')
    tree=ast.parse(source)
    names=['_excel_safe_value','_excel_safe_df','_write_df_to_ws_fast_h84']
    funcs=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in names]
    assert len(funcs)==3
    start=source.index('    sheets = [\n        ("00_主管重點觀察榜", h144_focus,')
    end=source.index('    return output.getvalue()',start)+len('    return output.getvalue()')
    wrapper='''def render(board,meta,authority,review):
    from openpyxl import Workbook
    import io
    wb=Workbook()
    wb.remove(wb.active)
    output=io.BytesIO()
    h144_focus=board
    h144_focus_meta=meta
    tables={'h126_authority':authority}
    h147_five_gates=review
    h147_five_meta={'reviewed':len(review),'data_gaps':0,'already_ignited':0}
    _h147_scan_diag={'status':'SCHEMA_RECONCILED'}
    _h145_bridge={'status':'SUPPLIED_FULL_SCAN_MATCHED'}
'''+source[start:end]+'\n'
    ns={'pd':pd}
    exec(compile(ast.fix_missing_locations(ast.Module(body=funcs,type_ignores=[])),str(PAGE),'exec'),ns,ns)
    exec(compile(wrapper,str(PAGE),'exec'),ns,ns)
    return ns['render']


def test_h147_actual_two_sheet_xlsx_has_five_factors_and_no_false_authority():
    from openpyxl import load_workbook
    raw={'股票代號':'2317','股票名稱':'測試股','類別':'AI',
         'H110最終研究優先分':52,'H110研究資訊截止日':'2026-10-28',
         'H114K線證據日':'2026-10-28','H113當日成交額百萬':200,
         'H113二十日均成交額百萬':160,'H113當日成交量張':1000,
         'H114K線量比':1.5,'H113流動性合格':'是',
         'H105發動階段':'ALREADY-IGNITED','H108族群生命週期':'FADE',
         'H109決策權威':'FADE_VETO','H123法人實證狀態':'PENDING_EVIDENCE',
         'H124大戶實證狀態':'PENDING_EVIDENCE'}
    audit=pd.DataFrame([raw]);authority=pd.DataFrame([{'股票名稱':'本輪無正式合格','H126能否買進':'禁止買進'}])
    state={'status':'EXPORT_SCOPE_INCOMPLETE','scope_incomplete':True,
           'scan_count':100,'diagnosis_count':20,'reported_final':1,'export_count':20}
    board,meta=build_focus_board({'audit':audit,'h126_authority':authority},state)
    diag,_=build_five_gate_review(audit)
    blob=_writer()(board,meta,authority,diag)
    wb=load_workbook(io.BytesIO(blob))
    assert wb.sheetnames==['00_主管重點觀察榜','00_H126主管決策']
    ws=wb.worksheets[0]
    assert ws['R7'].value=='H147｜成交量能'
    assert ws['S7'].value=='H147｜起漲位置'
    assert ws['T7'].value=='H147｜強勢族群'
    assert ws['U7'].value=='H147｜大戶鎖碼'
    assert ws['V7'].value=='H147｜技術結構'
    assert ws['R8'].value=='量能支持'
    assert ws['S8'].value=='已發動／防追高'
    assert ws['T8'].value=='退潮否決／待查來源'
    assert ws['U8'].value=='待補證'
    assert ws['M8'].value=='未列入H126合格決策'
    assert ws['P8'].value=='禁止依本榜買進'
    assert ws['B2'].value.endswith('SCHEMA_RECONCILED')
    assert ws.auto_filter.ref=='A7:Q8' and ws.freeze_panes=='D8'
    assert '本輪無正式合格' in [c.value for c in wb.worksheets[1][2]]
    wb.close()
