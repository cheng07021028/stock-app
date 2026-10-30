# -*- coding: utf-8 -*-
"""H148.3 adversarial research recall tests: evidence first, no trade authority."""
import io
from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook

from godpick_h1483_research_recall import build_research_recall, append_research_to_sheet

DAY='2026-10-30'


def sample(code='3035', **extra):
    row = {
       '股票代號':code,'股票名稱':'測試強勢股','類別':'半導體',
       'H114K線證據日':DAY,'H114K線資料狀態':'VALID_OHLCV',
       'H114K線有效根數':80,'H114K線量比':1.75,
       'H114當日成交量張':2600,'成交額百萬':375,
       '20日均成交額百萬':260,'H114距20日高點%':-0.8,
       '近5日漲幅%':6.3,'收盤距MA20%':4.0,'收盤距MA60%':8.2,
       'H109決策權威':'RESEARCH_ELIGIBLE','H123法人實證狀態':'PENDING_EVIDENCE',
       'H124大戶實證狀態':'PENDING_EVIDENCE',
    }
    row.update(extra)
    return row


def recall(candidates, audit=(), bridge=None):
    return build_research_recall(
         diagnosis_frame=pd.DataFrame(candidates),audit_frame=pd.DataFrame(audit),
         research_day=DAY,h145_bridge=bridge or {},limit=12)


def test_full_diagnosis_stock_not_in_h113_120_is_visible_research_not_trade():
    board,summary,meta=recall([sample('3035')],audit=[sample('2317')])
    assert len(board)==1 and board.iloc[0]['股票代號']=='3035'
    assert board.iloc[0]['研究層級'].startswith('A｜')
    assert '未進H113' in board.iloc[0]['尚需驗證與不得買進原因']
    assert 'H145同輪母體尚未確認' in board.iloc[0]['尚需驗證與不得買進原因']
    assert '非正式買進' in board.iloc[0]['尚需驗證與不得買進原因']
    assert not meta['formal_h126_modified'] and meta['scope_proven'] is False


def test_same_day_missing_average_is_B_not_fake_pass():
    board,_,meta=recall([sample('3035', **{'20日均成交額百萬':None})])
    assert len(board)==1 and board.iloc[0]['研究層級'].startswith('B｜')
    assert '20日均成交額未證' in board.iloc[0]['尚需驗證與不得買進原因']


def test_stale_day_is_C_never_current_strong_and_future_not_promoted():
    rows=[sample('3035',H114K線證據日='2026-10-29'),sample('2352',H114K線證據日='2026-11-03')]
    board,_,_=recall(rows)
    assert len(board)==2
    assert board['研究層級'].str.startswith('C｜').all()
    assert '行情未與本輪研究日一致' in board.iloc[0]['尚需驗證與不得買進原因']
    assert '同日' not in board.iloc[0]['研究層級']


def test_missing_ratio_never_invents_volume_breakout():
    board,summary,meta=recall([sample('3035',H114K線量比=None)])
    assert board.empty and meta['found']==0
    assert '缺乏' in summary.loc[summary['檢查項目']=='其他未列強勢線索','說明'].iloc[0] or meta['found']==0


def test_illiquid_high_gainer_never_fills_twelve():
    rows=[sample(str(2000+i),成交額百萬=35,**{'近5日漲幅%':12}) for i in range(12)]
    board,_,meta=recall(rows)
    assert board.empty and meta['found']==0


def test_score_only_and_52_ties_not_prioritized():
    rows=[sample('3035',H114K線量比=None, H110最終研究優先分=99),
          sample('2352',H110最終研究優先分=52)]
    board,_,_=recall(rows)
    assert list(board['股票代號'])==['2352']


def test_h109_fade_is_not_unvetoed_even_if_price_strong():
    board,_,_=recall([sample('2352',H109決策權威='FADE_VETO')])
    assert len(board)==1
    assert board.iloc[0]['H109風控']=='FADE_VETO'
    assert '原否決保留' in board.iloc[0]['尚需驗證與不得買進原因']


def test_repeated_code_conflict_quarantined_not_pick_better():
    rows=[sample('3035'),sample('3035',股票名稱='另一支'),sample('2352')]
    board,_,meta=recall(rows)
    assert list(board['股票代號'])==['2352']
    assert meta['duplicate_codes']==['3035']


def test_invalid_symbol_excluded_and_counted():
    board,_,meta=recall([sample('not_stock'),sample('3035')])
    assert len(board)==1 and meta['invalid_codes']==1


def test_no_candidate_leads_never_create_fake_stocks():
    board,summary,meta=recall([])
    assert board.empty and len(summary)>0
    assert meta['displayed']==0
    assert '不能' in meta['warning']


def test_sort_uses_actual_volume_ratio_not_stock_code_or_h110():
    rows=[sample('2001',H114K線量比=1.3),sample('9001',H114K線量比=2.0)]
    board,_,_=recall(rows)
    assert list(board['股票代號'])==['9001','2001']


def test_tier_A_before_B_C_and_limited_12():
    rows=[]
    for i in range(18):
        r=sample(str(3000+i))
        if i%3==1:r['20日均成交額百萬']=None
        elif i%3==2:r['H114K線證據日']='2026-10-29'
        rows.append(r)
    board,_,meta=recall(rows)
    assert len(board)==12 and meta['found']==18
    assert board.iloc[0]['研究層級'].startswith('A')
    assert list(board['研究層級'].str[0])==sorted(board['研究層級'].str[0])


def test_excel_same_two_sheets_and_append_cannot_modify_H126_or_filter():
    board,summary,_=recall([sample()])
    wb=Workbook(); ws=wb.active;ws.title='00_主管重點觀察榜'
    other=wb.create_sheet('00_H126主管決策')
    other['A1']='權威正式不可買';other['B2']='否'
    ws['A7']='觀察順位';ws['B7']='股票代號';ws.auto_filter.ref='A7:AJ7'
    res=append_research_to_sheet(ws,board,summary)
    assert res['research_rows']==1 and '研究線索' in ws.cell(res['first_row'],1).value
    assert ws.auto_filter.ref=='A7:AJ7' and other['B2'].value=='否'
    buf=io.BytesIO();wb.save(buf);buf.seek(0)
    read=load_workbook(buf)
    assert read.sheetnames==['00_主管重點觀察榜','00_H126主管決策']
    cells=[str(c.value or '') for row in read.worksheets[0] for c in row]
    assert '3035' in cells and any('非正式買進' in v for v in cells)
    assert read.worksheets[1]['B2'].value=='否'


def test_export_formula_injection_escaped():
    board,summary,_=recall([sample(股票名稱='=1+1')])
    wb=Workbook();ws=wb.active
    loc=append_research_to_sheet(ws,board,summary)
    name_cells=[ws.cell(i,3) for i in range(loc['first_row'],loc['last_row']+1)]
    assert any(str(cell.value).startswith("'=1+1") for cell in name_cells)
    assert all(cell.data_type!='f' for cell in name_cells)


def test_page07_integrates_diagnosis_not_limited_audit_and_ui():
    source=(Path(__file__).resolve().parents[1]/'pages'/'7_股神推薦.py').read_text('utf-8')
    assert 'diagnosis_frame=_h1483_source' in source
    assert 'audit_frame=tables.get(\'audit\', pd.DataFrame())' in source
    assert "h1483_named_research" in source
    assert 'append_research_to_sheet as _h1483_append' in source
    assert 'StockPick_H1483_ResearchRecall_' in source
    compile(source,'pages/7_股神推薦.py','exec')
