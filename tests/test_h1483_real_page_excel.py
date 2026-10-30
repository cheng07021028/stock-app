# -*- coding: utf-8 -*-
"""Execute the ACTUAL Page07 two-sheet export body with an H148.3 candidate."""
import ast
import io
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from godpick_h1483_research_recall import build_research_recall
from test_h1483_research_recall import sample,DAY


def test_actual_page07_exports_named_non_trade_research_even_with_no_h1481_strict_leads():
    src_path=Path(__file__).resolve().parents[1]/'pages'/'7_股神推薦.py'
    source=src_path.read_text('utf-8')
    tree=ast.parse(source)
    needed={'_excel_safe_value','_excel_safe_df','_write_df_to_ws_fast_h84'}
    nodes=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in needed]
    assert len(nodes)==3
    start=source.index('    sheets = [\n        ("00_主管重點觀察榜", h144_focus,')
    end=source.index('    return output.getvalue()', start)+len('    return output.getvalue()')
    block='''def export_page():
    import io
    from openpyxl import Workbook
    wb=Workbook();wb.remove(wb.active)
    output=io.BytesIO()
    h144_focus=pd.DataFrame(columns=['觀察順位','股票代號','股票名稱','買進資格'])
    h144_focus_meta={'scanned':2000,'diagnosed':1800,'final':0,'received':0,
      'audit_count':120,'authority_count':1,'all_scores_tied':False,
      'scope_issues':['H145母體未證'],'valid_score_rows':0,'score_distinct':0,
      'h126_buyable':0,'warning':'僅供研究'}
    tables={'h126_authority':pd.DataFrame([{'股票代號':'3035','H126能否買進':'否'}])}
    h147_five_gates=pd.DataFrame()
    h148_trace=pd.DataFrame()
    h148_meta={'review_unique':0,'scope_status':'母體未證','zero_formal_diagnosis':'0正式推薦'}
    h1481_meta={'eligible':0,'quarantine':120,'h113_missed_strong_research_count':5}
    h1483_research, h1483_summary, h1483_meta=build_research_recall(
       diagnosis_frame=pd.DataFrame([sample()]), audit_frame=pd.DataFrame(),
       research_day=DAY,h145_bridge={})
    h1482_repair=pd.DataFrame([{'檢查項目':'H113稽核','結果':120,'說明':'限量樣本'}])
    h1482_details=pd.DataFrame()
    h1482_meta={'same_day_h114':1,'candidate_not_h113':1}
    _h147_scan_diag={'status':'CHECK'}
    _h145_bridge={'status':'UNPROVEN_SOURCE'}
'''+source[start:end]+'\n'
    env={'pd':pd,'build_research_recall':build_research_recall,'sample':sample,'DAY':DAY}
    exec(compile(ast.fix_missing_locations(ast.Module(body=nodes,type_ignores=[])),str(src_path),'exec'),env,env)
    exec(compile(block,str(src_path),'exec'),env,env)
    wb=load_workbook(io.BytesIO(env['export_page']()))
    assert wb.sheetnames==['00_主管重點觀察榜','00_H126主管決策']
    ws,formal=wb.worksheets
    assert ws['B8'].value is None  # strict H148.1 board not falsified
    assert ws.auto_filter.ref.startswith('A7:')
    assert 'H148.3' in ws['A1'].value
    assert any(ws.cell(r,1).value and 'H148.3｜具名' in str(ws.cell(r,1).value)
               for r in range(8,ws.max_row+1))
    assert any(ws.cell(r,2).value=='3035' for r in range(8,ws.max_row+1))
    assert formal['B2'].value=='否'
    wb.close()
